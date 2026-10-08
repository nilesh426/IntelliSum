"""Output service — PDF and DOCX export of the structured summary (Phase 3).

Both builders consume the same structured payload produced by
``summarization_service.build_download_payload``::

    {
        "document_info": {...},
        "summary_statistics": {...},
        "executive_summary": "...",
        "key_points": [...],
        "section_summaries": [{"title": ..., "summary": ...}, ...],
        "keywords": [...],
        "important_concepts": [...],
        "important_sentences": [{"sentence": ..., "score": ...}, ...],
        "conclusion": "...",
    }

This guarantees the webpage, the PDF and the DOCX always show identical
numbers and text. Sections without content are omitted, never invented.
"""

import os
import re
from xml.sax.saxutils import escape

from werkzeug.utils import secure_filename

VALID_LENGTHS = ("short", "medium", "detailed")

METHODOLOGY_STEPS = [
    "Text extraction from the uploaded document.",
    "NLP preprocessing (normalization, tokenization, stopword handling).",
    "TF-IDF analysis to determine important terms.",
    "Sentence scoring with cosine similarity to the document centroid.",
    "Redundancy removal of near-duplicate sentences.",
    "Important sentence selection (highest-ranked sentences).",
    "Original sentence ordering for readability.",
]


def sanitize_stem(filename):
    """Return a filesystem-safe base name (no directories, no traversal)."""
    stem = os.path.splitext(os.path.basename(filename or "document"))[0]
    stem = secure_filename(stem.strip()) or "document"
    return stem


def export_filename(original_filename, length="medium", ext="pdf"):
    """Build a safe download filename like ``paper_summary_medium.pdf``."""
    length = (length or "medium").lower()
    if length not in VALID_LENGTHS:
        length = "medium"
    ext = re.sub(r"[^a-z0-9]", "", (ext or "pdf").lower()) or "pdf"
    return f"{sanitize_stem(original_filename)}_summary_{length}.{ext}"


def safe_output_path(output_dir, filename):
    """Join a filename to the output directory, blocking path escape."""
    os.makedirs(output_dir, exist_ok=True)
    safe_name = secure_filename(os.path.basename(filename or "summary.pdf"))
    if not safe_name:
        safe_name = "summary.pdf"
    dest = os.path.abspath(os.path.join(output_dir, safe_name))
    if os.path.dirname(dest) != os.path.abspath(output_dir):
        raise ValueError("Unsafe output filename.")
    return dest


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

def _stats_rows(payload):
    stats = payload.get("summary_statistics", {}) or {}
    doc_info = payload.get("document_info", {}) or {}
    page_count = doc_info.get("page_count")
    return [
        ("Summary Method", stats.get("method", "Extractive NLP")),
        ("Summary Length", stats.get("length", "")),
        ("Original Words", str(stats.get("original_words", 0))),
        ("Summary Words", str(stats.get("summary_words", 0))),
        (
            "Compression Ratio",
            f"{stats.get('compression_percent', 0)}% "
            f"({stats.get('summary_words', 0)} of "
            f"{stats.get('original_words', 0)} words)",
        ),
        ("Original Sentences", str(stats.get("original_sentences", 0))),
        ("Selected Sentences", str(stats.get("summary_sentences", 0))),
        ("Sections Detected", str(stats.get("section_count", 0))),
        ("Keywords Extracted", str(stats.get("keyword_count", 0))),
        (
            "Pages",
            str(page_count) if page_count is not None else "N/A",
        ),
    ]


def _doc_info_rows(payload):
    doc_info = payload.get("document_info", {}) or {}
    file_type = (doc_info.get("file_type", "") or "").upper() or "N/A"
    page_count = doc_info.get("page_count")
    if doc_info.get("file_type") == "pdf":
        pages = str(page_count) if page_count is not None else "N/A"
    else:
        pages = f"N/A ({file_type})" if file_type != "N/A" else "N/A"
    return [
        ("File Name", str(doc_info.get("filename", "N/A"))),
        ("File Type", file_type),
        ("Page Count", pages),
        ("Original Word Count", str(doc_info.get("word_count", 0))),
    ]


def _important_sentence_texts(payload):
    items = payload.get("important_sentences", []) or []
    texts = []
    for item in items:
        if isinstance(item, dict):
            sent = (item.get("sentence") or "").strip()
        else:
            sent = str(item).strip()
        if sent:
            texts.append(sent)
    return texts


# ---------------------------------------------------------------------------
# PDF export (ReportLab Platypus)
# ---------------------------------------------------------------------------

def build_pdf(payload, output_path):
    """Generate a paginated PDF of the structured summary.

    Uses ReportLab Platypus flowables so long summaries paginate
    cleanly (no overflow, no cut-off text, no manual positioning).
    Returns the absolute path of the written file.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    output_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "IntelliSumTitle", parent=styles["Title"], fontSize=22, spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        "IntelliSumSubtitle",
        parent=styles["Normal"],
        fontSize=11,
        textColor=colors.HexColor("#4b5563"),
        alignment=1,
        spaceAfter=10,
    )
    h1 = ParagraphStyle(
        "IntelliSumH1",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#111827"),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "IntelliSumH2",
        parent=styles["Heading2"],
        fontSize=12,
        textColor=colors.HexColor("#1f2937"),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )
    body = ParagraphStyle(
        "IntelliSumBody", parent=styles["Normal"], fontSize=10, leading=15,
    )
    cell_style = ParagraphStyle(
        "IntelliSumCell", parent=body, fontSize=10, leading=14,
    )
    cell_bold = ParagraphStyle(
        "IntelliSumCellBold", parent=cell_style, fontName="Helvetica-Bold",
    )

    def _page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawCentredString(A4[0] / 2.0, 1.5 * cm, f"Page {doc.page}")
        canvas.restoreState()

    story = [
        Paragraph("INTELLISUM", title_style),
        Paragraph(
            "Intelligent Document Summarization and Structuring System",
            subtitle_style,
        ),
    ]

    def _add_table(rows):
        data = [
            [Paragraph(f"<b>{escape(str(k))}</b>", cell_bold),
             Paragraph(escape(str(v)), cell_style)]
            for k, v in rows
        ]
        table = Table(data, colWidths=[5.2 * cm, 10.8 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f3f4f6")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
                ]
            )
        )
        story.append(table)

    # --- Document information ---
    story.append(Paragraph("DOCUMENT INFORMATION", h1))
    _add_table(_doc_info_rows(payload))

    # --- Summary overview ---
    story.append(Paragraph("SUMMARY OVERVIEW", h1))
    _add_table(_stats_rows(payload))

    # --- Executive summary ---
    executive = (payload.get("executive_summary") or "").strip()
    if executive:
        story.append(Paragraph("EXECUTIVE SUMMARY", h1))
        story.append(Paragraph(escape(executive), body))

    # --- Key points ---
    key_points = [p for p in (payload.get("key_points", []) or []) if str(p).strip()]
    if key_points:
        story.append(Paragraph("KEY POINTS", h1))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(escape(str(p)), body)) for p in key_points],
                bulletType="bullet",
                leftIndent=18,
            )
        )

    # --- Section-wise summary ---
    sections = [s for s in (payload.get("section_summaries", []) or [])
                if (s.get("summary") or "").strip()]
    if sections:
        story.append(Paragraph("SECTION-WISE SUMMARY", h1))
        for i, sec in enumerate(sections, start=1):
            story.append(Paragraph(f"{i}. {escape(str(sec.get('title', '')))}", h2))
            story.append(Paragraph(escape(str(sec.get("summary", ""))), body))

    # --- Keywords ---
    keywords = [k for k in (payload.get("keywords", []) or []) if str(k).strip()]
    if keywords:
        story.append(Paragraph("IMPORTANT KEYWORDS", h1))
        story.append(Paragraph(escape(" | ".join(keywords)), body))

    # --- Concepts ---
    concepts = [c for c in (payload.get("important_concepts", []) or [])
                if str(c).strip()]
    if concepts:
        story.append(Paragraph("IMPORTANT CONCEPTS", h1))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(escape(str(c)), body)) for c in concepts],
                bulletType="bullet",
                leftIndent=18,
            )
        )

    # --- Important sentences ---
    sentences = _important_sentence_texts(payload)
    if sentences:
        story.append(Paragraph("IMPORTANT SENTENCES", h1))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(escape(s), body)) for s in sentences],
                bulletType="1",
                leftIndent=24,
            )
        )

    # --- Conclusion ---
    conclusion = (payload.get("conclusion") or "").strip()
    if conclusion:
        story.append(Paragraph("CONCLUSION", h1))
        story.append(Paragraph(escape(conclusion), body))

    # --- Methodology ---
    story.append(Paragraph("HOW THE SUMMARY WAS GENERATED", h1))
    story.append(
        ListFlowable(
            [ListItem(Paragraph(escape(step), body)) for step in METHODOLOGY_STEPS],
            bulletType="1",
            leftIndent=24,
        )
    )
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            "Generated locally with extractive NLP (TF-IDF + cosine "
            "similarity). No external AI service was used.",
            ParagraphStyle(
                "IntelliSumNote", parent=body, fontSize=9,
                textColor=colors.HexColor("#6b7280"),
            ),
        )
    )

    template = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title="IntelliSum Summary",
        author="IntelliSum",
    )
    template.build(story, onFirstPage=_page_number, onLaterPages=_page_number)
    return output_path


# ---------------------------------------------------------------------------
# DOCX export (python-docx)
# ---------------------------------------------------------------------------

def build_docx(payload, output_path):
    """Generate a structured DOCX of the summary with proper heading levels.

    Returns the absolute path of the written file.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    output_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    doc = Document()
    doc.core_properties.title = "IntelliSum Summary"
    doc.core_properties.subject = (
        "Extractive NLP document summary (TF-IDF + cosine similarity)"
    )

    title = doc.add_heading("INTELLISUM", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Intelligent Document Summarization and Structuring System"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_table(rows):
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Light Grid Accent 1"
        for i, (key, value) in enumerate(rows):
            table.cell(i, 0).text = str(key)
            table.cell(i, 1).text = str(value)

    doc.add_heading("Document Information", level=1)
    _add_table(_doc_info_rows(payload))

    doc.add_heading("Summary Overview", level=1)
    _add_table(_stats_rows(payload))

    executive = (payload.get("executive_summary") or "").strip()
    if executive:
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(executive)

    key_points = [p for p in (payload.get("key_points", []) or []) if str(p).strip()]
    if key_points:
        doc.add_heading("Key Points", level=1)
        for point in key_points:
            doc.add_paragraph(str(point), style="List Bullet")

    sections = [s for s in (payload.get("section_summaries", []) or [])
                if (s.get("summary") or "").strip()]
    if sections:
        doc.add_heading("Section-wise Summary", level=1)
        for i, sec in enumerate(sections, start=1):
            doc.add_heading(f"{i}. {sec.get('title', '')}", level=2)
            doc.add_paragraph(str(sec.get("summary", "")))

    keywords = [k for k in (payload.get("keywords", []) or []) if str(k).strip()]
    if keywords:
        doc.add_heading("Important Keywords", level=1)
        doc.add_paragraph(" | ".join(keywords))

    concepts = [c for c in (payload.get("important_concepts", []) or [])
                if str(c).strip()]
    if concepts:
        doc.add_heading("Important Concepts", level=1)
        for concept in concepts:
            doc.add_paragraph(str(concept), style="List Bullet")

    sentences = _important_sentence_texts(payload)
    if sentences:
        doc.add_heading("Important Sentences", level=1)
        for sentence in sentences:
            doc.add_paragraph(sentence, style="List Number")

    conclusion = (payload.get("conclusion") or "").strip()
    if conclusion:
        doc.add_heading("Conclusion", level=1)
        doc.add_paragraph(conclusion)

    doc.add_heading("How the Summary Was Generated", level=1)
    for step in METHODOLOGY_STEPS:
        doc.add_paragraph(step, style="List Number")
    doc.add_paragraph(
        "Generated locally with extractive NLP (TF-IDF + cosine "
        "similarity). No external AI service was used."
    )

    doc.save(output_path)
    return output_path
