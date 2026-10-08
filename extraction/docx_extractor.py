"""DOCX text extraction using python-docx."""

import os

from docx import Document


def extract_docx(file_path):
    """Extract paragraphs, headings and table text from a DOCX file.

    Returns a consistent dict:
        {
            "filename": str,
            "file_type": "docx",
            "page_count": None,   # DOCX has no fixed pages
            "text": str,
            "metadata": dict,     # paragraph/headings/table info
            "headings": list[str],
        }

    Raises:
        FileNotFoundError: if the file does not exist.
        ValueError: if the file is corrupted/unreadable.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    filename = os.path.basename(file_path)

    try:
        doc = Document(file_path)
    except Exception as exc:
        raise ValueError(
            f"Could not read DOCX file '{filename}'. It may be corrupted."
        ) from exc

    try:
        paragraphs = []
        headings = []
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            style_name = getattr(para.style, "name", "") or ""
            if "Heading" in style_name:
                headings.append(text)
            paragraphs.append(text)

        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                row_text = " | ".join(c for c in cells if c)
                if row_text:
                    table_texts.append(row_text)

        parts = list(paragraphs)
        if table_texts:
            parts.append("\n".join(table_texts))
        text = "\n\n".join(parts).strip()

        # Core properties (title/author) where available.
        core = doc.core_properties
        metadata = {
            "paragraph_count": len(paragraphs),
            "heading_count": len(headings),
            "table_count": len(doc.tables),
            "title": (core.title or "") if core is not None else "",
            "author": (core.author or "") if core is not None else "",
            "subject": (core.subject or "") if core is not None else "",
        }
        metadata = {k: v for k, v in metadata.items() if v != "" and v is not None}

        return {
            "filename": filename,
            "file_type": "docx",
            "page_count": None,
            "text": text,
            "metadata": metadata,
            "headings": headings,
        }
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Failed to extract text from DOCX '{filename}'.") from exc
