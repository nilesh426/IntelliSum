"""Route tests: upload validation, supported/unsupported/empty/corrupt files."""

import io
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app import create_app  # noqa: E402


class TestRoutes(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = create_app()
        self.app.config["TESTING"] = True
        self.app.config["UPLOAD_FOLDER"] = self.tmp.name
        self.client = self.app.test_client()

    def tearDown(self):
        self.tmp.cleanup()

    def _post(self, filename, content):
        return self.client.post(
            "/summarize",
            data={"file": (io.BytesIO(content), filename)},
            content_type="multipart/form-data",
        )

    def test_index_loads(self):
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"IntelliSum", resp.data)
        self.assertIn(b"Summary length", resp.data)
        self.assertIn(b"Summarization method", resp.data)

    def test_txt_upload(self):
        text = (
            "Natural language processing enables computers to understand text. "
            "Machine learning learns patterns from large collections of documents. "
            "TF-IDF highlights the most important terms in each sentence. "
            "Cosine similarity compares sentence vectors for ranking purposes. "
            "Extractive methods select the most representative sentences directly."
        )
        resp = self._post("sample.txt", text.encode("utf-8"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Document Analysis", resp.data)

    def test_docx_upload(self):
        from docx import Document

        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph(
            "Natural language processing enables computers to understand text. " * 4
        )
        doc.add_heading("Methods", level=1)
        doc.add_paragraph("TF-IDF weighting highlights important terms. " * 4)
        buf = io.BytesIO()
        doc.save(buf)
        resp = self.client.post(
            "/summarize",
            data={"file": (io.BytesIO(buf.getvalue()), "sample.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Document Analysis", resp.data)

    def test_pdf_upload(self):
        try:
            import pymupdf
        except ImportError:
            import fitz as pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text(
            (72, 72),
            "Machine learning enables computers to learn from data. " * 10,
        )
        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        resp = self.client.post(
            "/summarize",
            data={"file": (io.BytesIO(buf.getvalue()), "sample.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Document Analysis", resp.data)

    def test_invalid_extension(self):
        resp = self._post("evil.exe", b"malicious content")
        self.assertEqual(resp.status_code, 302)  # redirect with flash message

    def test_empty_file(self):
        resp = self._post("empty.txt", b"")
        self.assertEqual(resp.status_code, 302)

    def test_corrupted_pdf(self):
        resp = self._post("broken.pdf", b"this is definitely not a pdf")
        self.assertIn(resp.status_code, (200, 302))

    def test_no_file_part(self):
        resp = self.client.post("/summarize", data={})
        self.assertEqual(resp.status_code, 302)


if __name__ == "__main__":
    unittest.main()
