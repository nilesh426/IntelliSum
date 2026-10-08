"""Tests for PDF / DOCX / TXT extraction (Phase 1)."""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from docx import Document  # noqa: E402

from extraction.docx_extractor import extract_docx  # noqa: E402
from extraction.pdf_extractor import extract_pdf  # noqa: E402
from extraction.txt_extractor import extract_txt  # noqa: E402


def _make_pdf(path, pages):
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()


class TestExtraction(unittest.TestCase):
    def test_txt_extraction(self):
        with tempfile.NamedTemporaryFile(
            suffix=".txt", delete=False, mode="w", encoding="utf-8"
        ) as fh:
            fh.write("Hello world.\nSecond line.")
            path = fh.name
        try:
            out = extract_txt(path)
            self.assertEqual(out["file_type"], "txt")
            self.assertIn("Hello world", out["text"])
            self.assertEqual(out["page_count"], 1)
        finally:
            os.remove(path)

    def test_txt_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            extract_txt("/nonexistent/file.txt")

    def test_docx_extraction(self):
        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This is a test paragraph about natural language processing.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Cell A"
        table.cell(0, 1).text = "Cell B"
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as fh:
            path = fh.name
        doc.save(path)
        try:
            out = extract_docx(path)
            self.assertEqual(out["file_type"], "docx")
            self.assertIn("natural language processing", out["text"])
            self.assertIn("Introduction", out["headings"])
            self.assertIn("Cell A", out["text"])
        finally:
            os.remove(path)

    def test_docx_corrupted(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, mode="w") as fh:
            fh.write("this is not a docx file {{{{")
            path = fh.name
        try:
            with self.assertRaises(ValueError):
                extract_docx(path)
        finally:
            os.remove(path)

    def test_pdf_extraction(self):
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as fh:
            path = fh.name
        _make_pdf(path, ["Machine learning is fascinating.", "Second page content here."])
        try:
            out = extract_pdf(path)
            self.assertEqual(out["file_type"], "pdf")
            self.assertEqual(out["page_count"], 2)
            self.assertIn("Machine learning", out["text"])
            self.assertEqual(len(out["pages"]), 2)
        finally:
            os.remove(path)

    def test_pdf_corrupted(self):
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False, mode="w") as fh:
            fh.write("not a pdf %%%%")
            path = fh.name
        try:
            with self.assertRaises(ValueError):
                extract_pdf(path)
        finally:
            os.remove(path)


if __name__ == "__main__":
    unittest.main()
